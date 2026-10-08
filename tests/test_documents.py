"""REQ-005 (DOCX → PDF), REQ-006 (Markdown/TXT → DOCX), REQ-007 (Markdown/TXT → PDF)."""

import io
import pathlib

import docx
import pymupdf
import pytest
from docx.oxml import parse_xml
from PIL import Image

from alihrupa.cli import main

NOTE = """# Judul Catatan

Paragraf dengan **tebal**, *miring*, dan huruf é ü ñ.

- satu
- dua

1. pertama

```
print("halo")
```
"""


def listing(path):
    return sorted(p.name for p in path.iterdir())


def pdf_text(path):
    with pymupdf.open(path) as doc:
        return "\n".join(page.get_text() for page in doc)


W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def make_docx(path, header=None, footer=None, textbox=None):
    d = docx.Document()
    d.styles["Normal"].font.name = "Libertinus Serif"  # bundled with Typst; Cambria would add a font warning
    d.add_heading("Dokumen Uji", 1)
    d.add_paragraph("Isi utama dokumen.")
    if header:
        d.sections[0].header.paragraphs[0].text = header
    if footer:
        d.sections[0].footer.paragraphs[0].text = footer
    if textbox:
        d.add_paragraph()._p.append(parse_xml(
            '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:shape style="width:100pt;height:50pt">'
            f"<v:textbox><w:txbxContent><w:p><w:r><w:t>{textbox}</w:t></w:r></w:p></w:txbxContent>"
            "</v:textbox></v:shape></w:pict></w:r>"
        ))
    d.save(path)
    return path


# --- REQ-006: Markdown/TXT → DOCX


def test_markdown_to_docx_keeps_formatting(tmp_path):
    source = tmp_path / "catatan.md"
    source.write_text(NOTE, encoding="utf-8")
    assert main([str(source), "docx"]) == 0
    d = docx.Document(tmp_path / "catatan.docx")
    by_text = {p.text: p for p in d.paragraphs}
    assert by_text["Judul Catatan"].style.name.startswith("Heading")
    para = by_text["Paragraf dengan tebal, miring, dan huruf é ü ñ."]
    assert any(r.bold and r.text == "tebal" for r in para.runs)
    assert any(r.italic and r.text == "miring" for r in para.runs)
    assert "satu" in by_text and "pertama" in by_text
    assert 'print("halo")' in by_text


def test_emoji_and_accents_survive(tmp_path):
    source = tmp_path / "emoji.md"
    source.write_text("Selamat 🎉 café\n", encoding="utf-8")
    assert main([str(source), "docx"]) == 0
    assert "Selamat 🎉 café" in [p.text for p in docx.Document(tmp_path / "emoji.docx").paragraphs]


def test_markdown_extension_alias(tmp_path):
    source = tmp_path / "catatan.markdown"
    source.write_text(NOTE, encoding="utf-8")
    assert main([str(source), "docx"]) == 0


def test_txt_is_not_read_as_markdown(tmp_path):
    source = tmp_path / "teks.txt"
    source.write_text("# bukan judul\nbaris dua **tetap bintang**\n\nparagraf baru", encoding="utf-8")
    assert main([str(source), "docx"]) == 0
    texts = [p.text for p in docx.Document(tmp_path / "teks.docx").paragraphs]
    assert "# bukan judul\nbaris dua **tetap bintang**" in texts
    assert "paragraf baru" in texts


@pytest.mark.parametrize("target", ["docx", "pdf"])
def test_invalid_utf8(tmp_path, capsys, target):
    source = tmp_path / "latin1.md"
    source.write_bytes("caf\xe9".encode("latin-1"))
    assert main([str(source), target]) == 1
    assert "UTF-8" in capsys.readouterr().err
    assert listing(tmp_path) == ["latin1.md"]


@pytest.mark.parametrize("target", ["docx", "pdf"])
def test_empty_file(tmp_path, target):
    source = tmp_path / "kosong.md"
    source.write_text("")
    assert main([str(source), target]) == 0
    assert (tmp_path / f"kosong.{target}").is_file()


@pytest.mark.parametrize("target", ["docx", "pdf"])
def test_remote_images_not_downloaded(tmp_path, capsys, http_server, target):
    base, requests = http_server
    source = tmp_path / "remote.md"
    source.write_text(f"Teks.\n\n![logo jauh]({base}/logo.png)\n\nSebaris ![ikon]({base}/i.png) di tengah.\n")
    assert main([str(source), target]) == 0
    assert requests == []
    assert "2 gambar dari URL tidak diunduh" in capsys.readouterr().err


@pytest.mark.parametrize("target", ["docx", "pdf"])
def test_local_relative_image_included(tmp_path, target):
    (tmp_path / "gambar").mkdir()
    Image.new("RGB", (60, 30), (0, 200, 0)).save(tmp_path / "gambar" / "hijau.png")
    source = tmp_path / "bergambar.md"
    source.write_text("Ada gambar:\n\n![hijau](gambar/hijau.png)\n")
    out = tmp_path / "hasil"
    assert main([str(source), target, "-o", str(out)]) == 0
    result = out / f"bergambar.{target}"
    if target == "docx":
        assert len(docx.Document(result).inline_shapes) == 1
    else:
        with pymupdf.open(result) as doc:
            assert len(doc[0].get_images()) == 1


# --- REQ-005: DOCX → PDF


def test_docx_to_pdf(tmp_path, capsys):
    source = make_docx(tmp_path / "surat.docx")
    assert main([str(source), "pdf"]) == 0
    text = pdf_text(tmp_path / "surat.pdf")
    assert text.index("Dokumen Uji") < text.index("Isi utama dokumen.")
    assert "peringatan" not in capsys.readouterr().err


def test_docx_dropped_parts_are_reported(tmp_path, capsys):
    source = make_docx(tmp_path / "kop.docx", header="KOP", footer="KAKI", textbox="KOTAK")
    assert main([str(source), "pdf"]) == 0
    err = capsys.readouterr().err
    assert "header" in err and "footer" in err and "text box" in err
    assert "Isi utama dokumen." in pdf_text(tmp_path / "kop.pdf")


def test_docx_rich_content_in_pdf(tmp_path):
    d = docx.Document()
    d.add_heading("Judul", 1)
    p = d.add_paragraph()
    p.add_run("TEBAL").bold = True
    p.add_run(" dan ")
    p.add_run("MIRING").italic = True
    d.add_paragraph("butir satu", style="List Bullet")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(1, 1).text = "SEL-A", "SEL-D"
    buf = io.BytesIO()
    Image.new("RGB", (60, 30), (0, 200, 0)).save(buf, "PNG")
    buf.seek(0)
    d.add_picture(buf)
    d.save(tmp_path / "kaya.docx")

    assert main([str(tmp_path / "kaya.docx"), "pdf"]) == 0
    with pymupdf.open(tmp_path / "kaya.pdf") as doc:
        page = doc[0]
        fonts = {s["text"]: s["font"] for b in page.get_text("dict")["blocks"]
                 for line in b.get("lines", []) for s in line["spans"]}
        assert "Bold" in fonts["TEBAL"] and "Italic" in fonts["MIRING"]
        assert {"butir satu", "SEL-A", "SEL-D"} <= fonts.keys()
        assert len(page.get_images()) == 1


def test_empty_header_is_not_reported(tmp_path, capsys):
    """Word's default header holds only tab stops; that is not lost content."""
    d = docx.Document()
    d.add_paragraph("Isi.")
    d.sections[0].header.paragraphs[0]._p.get_or_add_pPr().append(
        parse_xml(f'<w:tabs {W}><w:tab w:val="center" w:pos="4680"/></w:tabs>')
    )
    d.save(tmp_path / "header_kosong.docx")
    assert main([str(tmp_path / "header_kosong.docx"), "pdf"]) == 0
    assert "tidak ikut ke PDF" not in capsys.readouterr().err


def test_columns_are_reported(tmp_path, capsys):
    d = docx.Document()
    d.add_paragraph("Kolom.")
    d.sections[0]._sectPr.append(parse_xml(f'<w:cols {W} w:space="720" w:num="2"/>'))
    d.save(tmp_path / "kolom.docx")
    assert main([str(tmp_path / "kolom.docx"), "pdf"]) == 0
    assert "tata letak kolom" in capsys.readouterr().err


@pytest.mark.parametrize("content", [b"bukan zip", b"PK\x03\x04rusak"])
def test_corrupt_docx(tmp_path, capsys, content):
    source = tmp_path / "rusak.docx"
    source.write_bytes(content)
    assert main([str(source), "pdf"]) == 1
    assert "rusak.docx" in capsys.readouterr().err
    assert listing(tmp_path) == ["rusak.docx"]


def test_no_office_suite_dependency():
    """G-004: nothing in the code or dependencies refers to an office suite."""
    root = pathlib.Path(__file__).parent.parent
    sources = [*root.joinpath("src").rglob("*.py"), root / "pyproject.toml"]
    for path in sources:
        text = path.read_text().lower()
        for word in ("soffice", "libreoffice", "unoconv", "docx2pdf"):
            assert word not in text, f"{word} in {path}"


# --- REQ-007: Markdown/TXT → PDF


def test_markdown_to_pdf_without_intermediate_files(tmp_path):
    source = tmp_path / "catatan.md"
    source.write_text(NOTE, encoding="utf-8")
    assert main([str(source), "pdf"]) == 0
    text = pdf_text(tmp_path / "catatan.pdf")
    assert "Judul Catatan" in text and "é ü ñ" in text and 'print("halo")' in text
    assert listing(tmp_path) == ["catatan.md", "catatan.pdf"]


def test_existing_docx_next_to_source_untouched(tmp_path):
    source = tmp_path / "catatan.md"
    source.write_text(NOTE, encoding="utf-8")
    (tmp_path / "catatan.docx").write_bytes(b"punya pengguna")
    assert main([str(source), "pdf"]) == 0
    assert (tmp_path / "catatan.docx").read_bytes() == b"punya pengguna"


def test_txt_to_pdf(tmp_path):
    source = tmp_path / "teks.txt"
    source.write_text("# bukan judul\n\nparagraf kedua", encoding="utf-8")
    assert main([str(source), "pdf"]) == 0
    assert "# bukan judul" in pdf_text(tmp_path / "teks.pdf")
