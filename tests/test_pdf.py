"""REQ-004: PDF → DOCX."""

import io

import docx
import pymupdf
import pytest
from PIL import Image

from alihrupa.cli import main


def png_bytes(color="gray", size=(200, 100)):
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, "PNG")
    return buf.getvalue()


def text_pdf(path, with_image=False):
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 80), "Laporan Bulanan", fontsize=22)
    page.insert_text((72, 130), "Paragraf pertama dengan huruf é dan ü.", fontsize=11)
    page.insert_text((72, 400), "Paragraf terakhir.", fontsize=11)
    if with_image:
        page.insert_image(pymupdf.Rect(72, 200, 272, 300), stream=png_bytes("blue"))
    doc.save(path)
    return path


def docx_text(path):
    return "\n".join(p.text for p in docx.Document(path).paragraphs)


def test_text_pdf_becomes_editable_docx(tmp_path, capsys):
    source = text_pdf(tmp_path / "laporan.pdf")
    assert main([str(source), "docx"]) == 0
    text = docx_text(tmp_path / "laporan.docx")
    assert text.index("Laporan Bulanan") < text.index("Paragraf pertama") < text.index("Paragraf terakhir")
    assert "é dan ü" in text
    out, err = capsys.readouterr()
    # Only our own result line: no library logging leaking into the terminal.
    assert out.strip() == f"{source} → {tmp_path / 'laporan.docx'}"
    assert err == ""


def test_images_are_kept(tmp_path):
    source = text_pdf(tmp_path / "bergambar.pdf", with_image=True)
    assert main([str(source), "docx"]) == 0
    assert len(docx.Document(tmp_path / "bergambar.docx").inline_shapes) >= 1


def test_scanned_pdf_converts_with_warning(tmp_path, capsys):
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(page.rect, stream=png_bytes())
    doc.save(tmp_path / "scan.pdf")
    assert main([str(tmp_path / "scan.pdf"), "docx"]) == 0
    assert "OCR belum didukung" in capsys.readouterr().err
    assert (tmp_path / "scan.docx").is_file()


def test_encrypted_pdf(tmp_path, capsys):
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "rahasia")
    doc.save(tmp_path / "kunci.pdf", encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="pw", owner_pw="pw2")
    assert main([str(tmp_path / "kunci.pdf"), "docx"]) == 1
    assert "berpassword" in capsys.readouterr().err
    assert sorted(p.name for p in tmp_path.iterdir()) == ["kunci.pdf"]


@pytest.mark.parametrize("content", [b"%PDF-1.4 rusak", b"bukan pdf sama sekali"])
def test_corrupt_pdf(tmp_path, capsys, content):
    source = tmp_path / "rusak.pdf"
    source.write_bytes(content)
    assert main([str(source), "docx"]) == 1
    assert "rusak.pdf" in capsys.readouterr().err
    assert sorted(p.name for p in tmp_path.iterdir()) == ["rusak.pdf"]


def test_skipped_page_is_reported(tmp_path, capsys, monkeypatch):
    from pdf2docx.page.Page import Page

    original, calls = Page.parse, []

    def first_page_breaks(self, **kwargs):
        calls.append(1)
        if len(calls) == 1:
            raise ValueError("halaman aneh")
        return original(self, **kwargs)

    monkeypatch.setattr(Page, "parse", first_page_breaks)
    doc = pymupdf.open()
    for n in (1, 2):
        doc.new_page().insert_text((72, 72), f"Halaman {n}")
    doc.save(tmp_path / "dua.pdf")
    assert main([str(tmp_path / "dua.pdf"), "docx"]) == 0
    err = capsys.readouterr().err
    assert "dilewati" in err and "halaman aneh" in err
    assert "Halaman 2" in docx_text(tmp_path / "dua.docx")
