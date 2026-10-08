"""REQ-011: DOCX → PDF follows the heading styles and the first-line indent of the DOCX."""

import io
import re
import zipfile

import docx
import pymupdf
import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from alihrupa.cli import main
from test_docx_layout import PT_PER_MM, convert, lines, mm, page_mm

BODY = "Isi paragraf yang cukup panjang supaya membungkus ke beberapa baris di dalam halaman ini. " * 5
FONT = "Libertinus Serif"  # bundled with Typst, so no font warnings
LEFT = 25.4  # margin in every test, mm


def build(path, blocks, *, indent=None, size=12, line=1.0, before=0, after=6, styles=None, tweak=None):
    """A DOCX of ("h1", text) / ("h2", text) / ("p", text) / ("li", text) / ("table", text) blocks.

    styles: {1: {...}} properties of Heading 1..n: align, size, bold, italic, color, font, before, after,
    line, page_break.
    """
    d = docx.Document()
    normal = d.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(size)
    normal.paragraph_format.line_spacing = line
    normal.paragraph_format.space_before, normal.paragraph_format.space_after = Pt(before), Pt(after)
    if indent is not None:
        normal.paragraph_format.first_line_indent = Cm(indent)
    sect = d.sections[0]
    sect.page_width, sect.page_height = Cm(21), Cm(29.7)  # A4; python-docx's default is Letter
    sect.top_margin = sect.bottom_margin = sect.left_margin = sect.right_margin = Cm(2.54)
    for level in range(1, 10):
        style = d.styles[f"Heading {level}"]
        style.font.name = FONT
        for key, value in (styles or {}).get(level, {}).items():
            p, f = style.paragraph_format, style.font
            if key == "align":
                p.alignment = value
            elif key == "size":
                f.size = Pt(value)
            elif key == "bold":
                f.bold = value
            elif key == "italic":
                f.italic = value
            elif key == "color":
                f.color.rgb = RGBColor.from_string(value)
            elif key == "font":
                f.name = value
            elif key == "before":
                p.space_before = Pt(value)
            elif key == "after":
                p.space_after = Pt(value)
            elif key == "line":
                p.line_spacing = value
            elif key == "page_break":
                p.page_break_before = value
    for kind, text in blocks:
        if kind[0] == "h":
            d.add_heading(text, int(kind[1]))
        elif kind == "li":
            d.add_paragraph(text, style="List Bullet")
        elif kind == "table":
            cell = d.add_table(rows=1, cols=1).cell(0, 0)
            cell.text = text
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT  # pandoc would centre an unaligned cell
        else:
            d.add_paragraph(text)
    d.save(path)
    if tweak:
        rewrite(path, tweak)
    return path


def rewrite(path, tweak):
    """Apply tweak(styles_xml) -> styles_xml to a saved DOCX."""
    buf = io.BytesIO()
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "word/styles.xml":
                data = tweak(data.decode("utf-8")).encode("utf-8")
            out.writestr(item, data)
    path.write_bytes(buf.getvalue())


def style_xml(styles, level):
    return re.search(rf'<w:style [^>]*w:styleId="Heading{level}".*?</w:style>', styles, re.S).group(0)


def swap(level, old, new):
    """A tweak replacing old with new inside Heading <level>'s style."""
    return lambda styles: styles.replace(style_xml(styles, level), re.sub(old, new, style_xml(styles, level), count=1))


def line_of(pdf, text, page=None):
    with pymupdf.open(pdf) as doc:
        n = doc.page_count
    for p in range(n) if page is None else [page]:
        for l in lines(pdf, p):
            if l[4].startswith(text):
                return l
    raise AssertionError(f"{text!r} not found")


def colour(pdf, text):
    with pymupdf.open(pdf) as doc:
        for page in doc:
            for b in page.get_text("dict")["blocks"]:
                for line in b.get("lines", []):
                    for s in line["spans"]:
                        if s["text"].startswith(text):
                            return s["color"]


# AC-1: alignment, size, bold, italic per level
def test_alignment_size_and_style_per_level(tmp_path):
    pdf = convert(build(tmp_path / "a.docx", [("h1", "Judul Satu"), ("p", BODY), ("h2", "Judul Dua"), ("p", BODY),
                                              ("h3", "Judul Tiga"), ("p", BODY)],
                        styles={1: dict(align=WD_ALIGN_PARAGRAPH.CENTER, size=14, bold=True),
                                2: dict(align=WD_ALIGN_PARAGRAPH.LEFT, size=12, bold=True),
                                3: dict(align=WD_ALIGN_PARAGRAPH.RIGHT, size=11, bold=False, italic=True)}))
    h1, h2, h3 = (line_of(pdf, t) for t in ("Judul Satu", "Judul Dua", "Judul Tiga"))
    text_left, text_right = LEFT, 210 - LEFT
    assert (mm(h1[0]) + mm(h1[2])) / 2 == pytest.approx((text_left + text_right) / 2, abs=0.5)
    assert mm(h2[0]) == pytest.approx(text_left, abs=0.5)
    assert mm(h3[2]) == pytest.approx(text_right, abs=0.5)
    assert (h1[6], h2[6], h3[6]) == (14, 12, 11)
    assert "Bold" in h1[5] and "Bold" in h2[5]
    assert "Italic" in h3[5] and "Bold" not in h3[5]


# AC-2: colour and typeface
def test_heading_colour(tmp_path):
    pdf = convert(build(tmp_path / "a.docx", [("h1", "Judul Biru"), ("p", BODY)], styles={1: dict(color="1F3864")}))
    assert colour(pdf, "Judul Biru") == 0x1F3864
    assert colour(pdf, "Isi paragraf") == 0


def test_heading_font_installed(tmp_path):
    pdf = convert(build(tmp_path / "a.docx", [("h1", "Judul"), ("p", BODY)], styles={1: dict(font="New Computer Modern")}))
    assert "NewCM" in line_of(pdf, "Judul")[5]
    assert "Libertinus" in line_of(pdf, "Isi paragraf")[5]


def test_heading_font_missing_warns_and_uses_body_font(tmp_path, capsys):
    pdf = convert(build(tmp_path / "a.docx", [("h1", "Judul"), ("p", BODY)], styles={1: dict(font="Font Heading Hilang")}))
    err = capsys.readouterr().err
    assert "Font Heading Hilang" in err and "peringatan" in err
    assert "Libertinus" in line_of(pdf, "Judul")[5]


# AC-3: spacing around headings (Word adds the paragraph's space after to the heading's space before)
def test_heading_spacing(tmp_path):
    pdf = convert(build(tmp_path / "a.docx", [("p", "Sebelum."), ("h1", "Judul"), ("p", "Sesudah.")], size=12, line=1.0, after=6,
                        styles={1: dict(size=14, before=12, after=18, line=1.5)}))
    before, head, after = (line_of(pdf, t) for t in ("Sebelum.", "Judul", "Sesudah."))
    descent = 0.2
    assert head[3] - before[3] == pytest.approx(descent * 12 + (6 + 12) + (1.5 * 1.15 - descent) * 14, abs=1)
    assert after[3] - head[3] == pytest.approx(descent * 14 + (18 + 0) + (1.0 * 1.15 - descent) * 12, abs=1)


# AC-4: page break before
def test_page_break_before(tmp_path):
    pdf = convert(build(tmp_path / "a.docx", [("h1", "Bab Satu"), ("p", BODY), ("h1", "Bab Dua"), ("p", BODY),
                                              ("h2", "Tanpa Pemisah"), ("p", BODY)],
                        styles={1: dict(page_break=True)}))
    with pymupdf.open(pdf) as doc:
        assert doc.page_count == 2
        assert all(p.get_text().strip() for p in doc)
    assert lines(pdf, 0)[0][4] == "Bab Satu"
    assert lines(pdf, 1)[0][4] == "Bab Dua"
    assert line_of(pdf, "Tanpa Pemisah")[3] > lines(pdf, 1)[0][3]


def test_no_page_break_without_style(tmp_path):
    pdf = convert(build(tmp_path / "a.docx", [("p", BODY), ("h1", "Bab"), ("p", BODY)]))
    with pymupdf.open(pdf) as doc:
        assert doc.page_count == 1


# AC-5: first-line indent
def test_first_line_indent(tmp_path):
    pdf = convert(build(tmp_path / "a.docx", [("h1", "Judul"), ("p", BODY), ("p", BODY)], indent=1.25))
    first = line_of(pdf, "Isi paragraf")  # first paragraph after the heading is indented too
    assert mm(first[0]) == pytest.approx(LEFT + 12.5, abs=0.5)
    body = [l for l in lines(pdf) if l[4] != "Judul"]
    assert sum(1 for l in body if mm(l[0]) > LEFT + 10) == 2  # exactly the two paragraph starts
    assert all(mm(l[0]) == pytest.approx(LEFT, abs=0.5) for l in body if mm(l[0]) <= LEFT + 10)
    assert mm(line_of(pdf, "Judul")[0]) == pytest.approx(LEFT, abs=0.5)


def test_no_indent_in_tables_and_lists(tmp_path):
    pdf = convert(build(tmp_path / "a.docx", [("p", BODY), ("li", "Butir daftar"), ("table", "Isi sel")], indent=1.25))
    assert mm(line_of(pdf, "Isi sel")[0]) < LEFT + 4  # cell padding only
    assert mm(line_of(pdf, "•")[0]) == pytest.approx(LEFT, abs=0.5)


def test_no_indent_when_docx_has_none(tmp_path):
    pdf = convert(build(tmp_path / "a.docx", [("p", BODY), ("p", BODY)]))
    assert all(mm(l[0]) == pytest.approx(LEFT, abs=0.5) for l in lines(pdf))


def test_hanging_indent_is_not_a_first_line_indent(tmp_path):
    def hanging(styles):
        return styles.replace('<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/>',
                              '<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/>'
                              '<w:pPr><w:ind w:left="709" w:hanging="709"/></w:pPr>', 1)
    pdf = convert(build(tmp_path / "a.docx", [("p", BODY)], tweak=hanging))
    assert all(mm(l[0]) == pytest.approx(LEFT, abs=0.5) for l in lines(pdf))


# AC-6: styles or values that are not there
def test_missing_values_fall_back(tmp_path):
    def strip_sizes(styles):
        return re.sub(r"<w:szCs? [^>]*/>", "", styles)

    pdf = convert(build(tmp_path / "a.docx", [("h1", "Judul Tanpa Ukuran"), ("p", BODY)], tweak=strip_sizes))
    assert line_of(pdf, "Judul Tanpa Ukuran")[6] > 12  # Typst's own heading size, as before this REQ


# AC-7: nonsense values are ignored
@pytest.mark.parametrize("old,new", [
    (r'<w:sz w:val="\d+"/>', '<w:sz w:val="0"/>'),
    (r'<w:sz w:val="\d+"/>', '<w:sz w:val="2000"/>'),
    (r'<w:color w:val="\w+"[^>]*/>', '<w:color w:val="ZZZZZZ"/>'),
    (r'<w:spacing [^>]*/>', '<w:spacing w:before="-500" w:after="9999999"/>'),
])
def test_nonsense_heading_values_are_ignored(tmp_path, old, new):
    pdf = convert(build(tmp_path / "a.docx", [("h1", "Judul"), ("p", BODY)], tweak=swap(1, old, new)))
    head = line_of(pdf, "Judul")
    assert 6 <= head[6] <= 400 and head[6] != 1000
    assert colour(pdf, "Judul") == 0 or colour(pdf, "Judul") == 0x365F91
    with pymupdf.open(pdf) as doc:
        assert doc.page_count == 1


@pytest.mark.parametrize("value", ["-700", "99999"])
def test_nonsense_indent_is_ignored(tmp_path, value):
    def bad_indent(styles):
        return styles.replace('<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/>',
                              f'<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/>'
                              f'<w:pPr><w:ind w:firstLine="{value}"/></w:pPr>', 1)
    pdf = convert(build(tmp_path / "a.docx", [("p", BODY), ("p", BODY)], tweak=bad_indent))
    assert all(mm(l[0]) == pytest.approx(LEFT, abs=0.5) for l in lines(pdf))


def test_corrupt_docx_still_fails(tmp_path, capsys):
    src = tmp_path / "rusak.docx"
    src.write_bytes(b"PK\x03\x04rusak")
    assert main([str(src), "pdf"]) == 1
    assert sorted(p.name for p in tmp_path.iterdir()) == ["rusak.docx"]


# AC-8: nothing changes without heading styles or indent
def test_markdown_headings_unchanged(tmp_path):
    src = tmp_path / "c.md"
    src.write_text("# Judul\n\nIsi.\n\n## Sub\n\nIsi lagi.\n", encoding="utf-8")
    pdf = convert(src)
    head = line_of(pdf, "Judul")
    assert mm(head[0]) == pytest.approx(31.75, abs=0.5)
    assert "Libertinus" in head[5] and head[6] > 11
