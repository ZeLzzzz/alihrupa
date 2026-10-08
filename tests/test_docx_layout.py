"""REQ-010: DOCX → PDF follows the DOCX's basic layout (margins, font, spacing, alignment)."""

import docx
import pymupdf
import pytest
import typst
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import parse_xml
from docx.shared import Cm, Pt

from alihrupa.cli import main
from alihrupa.documents import check_conf
from alihrupa.errors import ConvertError

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
PT_PER_MM = 72 / 25.4
TEXT = ("Pemodelan proses menggambarkan struktur organisasi supaya bisa dipahami dianalisis dan dirancang ulang "
        "bila perlu oleh analis bisnis maupun manajer yang bertanggung jawab atas proses tersebut. ") * 4
ENGLISH = ("Comprehensibility of conceptual process documentation remains questionable whenever organizational "
           "representations incorporate extraordinarily complicated interdependencies. ") * 6
FONTS = {f.lower() for f in typst.Fonts().families()}


def make_docx(path, *, paragraphs=(TEXT,), font=None, size=None, line=None, before=None, after=None,
              align=None, page=None, margins=None, hyphenate=False, page_field=False, footer=None, title=None):
    """A DOCX with the given layout. line: a multiple (1.5) or Pt(...) for exact spacing.

    The font defaults to Libertinus Serif (bundled with Typst): python-docx's own default, Cambria,
    is often not installed and would add a font warning to every test.
    """
    d = docx.Document()
    normal = d.styles["Normal"]
    normal.font.name = font or "Libertinus Serif"
    if size:
        normal.font.size = Pt(size)
    fmt = normal.paragraph_format
    if line is not None:
        fmt.line_spacing = line
        if not isinstance(line, float):
            fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    if before is not None:
        fmt.space_before = Pt(before)
    if after is not None:
        fmt.space_after = Pt(after)
    if align is not None:
        fmt.alignment = align
    section = d.sections[0]
    if page:
        section.page_width, section.page_height = Cm(page[0]), Cm(page[1])
    if margins:
        section.top_margin, section.right_margin, section.bottom_margin, section.left_margin = (Cm(m) for m in margins)
    if hyphenate:
        d.settings.element.append(parse_xml(f"<w:autoHyphenation {W}/>"))
    if page_field:
        section.footer.paragraphs[0]._p.append(parse_xml(
            f'<w:fldSimple {W} w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>'))
    if footer:
        section.footer.paragraphs[0].text = footer
    if title:
        d.add_heading(title, 0)
    for text in paragraphs:
        d.add_paragraph(text)
    d.save(path)
    return path


def convert(src, *extra):
    assert main([str(src), "pdf", *extra]) == 0
    return src.with_suffix(".pdf")


def lines(pdf, page=0):
    """Text lines of a page: (x0, top, x1, baseline, text, font, size), top to bottom.

    x0/x1 come from the visible glyphs, so a trailing space does not count; a hyphen added by
    line breaking (extracted as a soft hyphen) is returned as "-".
    """
    with pymupdf.open(pdf) as doc:
        out = []
        for block in doc[page].get_text("rawdict")["blocks"]:
            for line in block.get("lines", []):
                spans = [s for s in line["spans"] if "".join(c["c"] for c in s["chars"]).strip()]
                chars = [c for s in spans for c in s["chars"] if c["c"].strip()]
                if chars:
                    text = "".join(c["c"] for s in spans for c in s["chars"]).replace("\xad", "-")
                    out.append((min(c["bbox"][0] for c in chars), line["bbox"][1], max(c["bbox"][2] for c in chars),
                                spans[0]["origin"][1], text, spans[0]["font"], spans[0]["size"]))
        return sorted(out, key=lambda l: l[3])


def page_mm(pdf):
    with pymupdf.open(pdf) as doc:
        r = doc[0].rect
        return r.width / PT_PER_MM, r.height / PT_PER_MM


def mm(pt):
    return pt / PT_PER_MM


def all_text(pdf):
    with pymupdf.open(pdf) as doc:
        return "\n".join(p.get_text() for p in doc)


# AC-1: margins and page size
def test_margins_and_page_size(tmp_path):
    src = make_docx(tmp_path / "a.docx", page=(21.5, 33), margins=(3, 2, 2.5, 4), size=12, line=1.0,
                    align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    pdf = convert(src)
    w, h = page_mm(pdf)
    assert w == pytest.approx(215, abs=1) and h == pytest.approx(330, abs=1)
    body = lines(pdf)
    assert mm(min(l[0] for l in body)) == pytest.approx(40, abs=0.5)  # left
    assert mm(max(l[2] for l in body)) == pytest.approx(215 - 20, abs=0.5)  # right (justified lines)
    # Like Word, the first line box starts at the top margin; its baseline sits one line height below, minus the descent.
    assert mm(body[0][3]) == pytest.approx(30 + mm(12 * (1.15 - 0.2)), abs=0.5)


def test_paper_option_wins_but_margins_stay(tmp_path):
    src = make_docx(tmp_path / "a.docx", page=(21.5, 33), margins=(3, 2, 2.5, 4))
    pdf = convert(src, "-p", "a5")
    w, h = page_mm(pdf)
    assert w == pytest.approx(148, abs=1) and h == pytest.approx(210, abs=1)
    assert mm(min(l[0] for l in lines(pdf))) == pytest.approx(40, abs=0.5)


def test_no_blank_first_page_and_title_kept(tmp_path):
    src = make_docx(tmp_path / "a.docx", title="Judul Dokumen", margins=(2.54,) * 4)
    pdf = convert(src)
    with pymupdf.open(pdf) as doc:
        assert all(p.get_text().strip() for p in doc)
    assert "Judul Dokumen" in all_text(pdf)


# AC-2: font size and family
def test_font_size(tmp_path):
    pdf = convert(make_docx(tmp_path / "a.docx", size=13))
    assert {round(l[6], 1) for l in lines(pdf)} == {13.0}


def test_font_installed_is_used(tmp_path):
    # New Computer Modern ships with Typst, so it is installed everywhere.
    pdf = convert(make_docx(tmp_path / "a.docx", font="New Computer Modern"))
    assert all("NewCM" in l[5] for l in lines(pdf))


@pytest.mark.skipif("liberation serif" not in FONTS, reason="Liberation Serif not installed")
def test_microsoft_font_gets_metric_compatible_substitute(tmp_path, capsys):
    pdf = convert(make_docx(tmp_path / "a.docx", font="Times New Roman"))
    if "times new roman" not in FONTS:
        assert all("LiberationSerif" in l[5] for l in lines(pdf))
    assert "tidak terpasang" not in capsys.readouterr().err


# AC-3: line and paragraph spacing
@pytest.mark.parametrize("multiple", [1.0, 1.5, 2.0])
def test_line_spacing_multiple(tmp_path, multiple):
    pdf = convert(make_docx(tmp_path / "a.docx", size=12, line=multiple, after=0))
    body = lines(pdf)
    gaps = [b[3] - a[3] for a, b in zip(body, body[1:])]
    assert len(gaps) >= 3
    for gap in gaps:
        assert gap == pytest.approx(1.15 * multiple * 12, rel=0.05)


def test_exact_line_spacing(tmp_path):
    pdf = convert(make_docx(tmp_path / "a.docx", size=12, line=Pt(20), after=0))
    body = lines(pdf)
    for a, b in zip(body, body[1:]):
        assert b[3] - a[3] == pytest.approx(20, rel=0.05)


def test_paragraph_spacing(tmp_path):
    pdf = convert(make_docx(tmp_path / "a.docx", paragraphs=["Satu.", "Dua.", "Tiga."], size=12, line=1.0,
                            before=6, after=10))
    body = lines(pdf)
    assert [l[4] for l in body] == ["Satu.", "Dua.", "Tiga."]
    for a, b in zip(body, body[1:]):
        assert b[3] - a[3] == pytest.approx(1.15 * 12 + 10 + 6, rel=0.05)


# AC-4: alignment
def test_justified_stays_justified(tmp_path):
    pdf = convert(make_docx(tmp_path / "a.docx", align=WD_ALIGN_PARAGRAPH.JUSTIFY, margins=(2.54,) * 4))
    full = lines(pdf)[:-1]  # the last line of a paragraph is never stretched
    assert max(l[2] for l in full) - min(l[2] for l in full) < 1


def test_left_stays_left(tmp_path):
    pdf = convert(make_docx(tmp_path / "a.docx", align=WD_ALIGN_PARAGRAPH.LEFT, margins=(2.54,) * 4))
    full = lines(pdf)[:-1]
    assert max(l[2] for l in full) - min(l[2] for l in full) > 5


def test_no_alignment_means_left(tmp_path):
    pdf = convert(make_docx(tmp_path / "a.docx", margins=(2.54,) * 4))
    full = lines(pdf)[:-1]
    assert max(l[2] for l in full) - min(l[2] for l in full) > 5


# AC-5: hyphenation and page numbers
def test_no_hyphenation_by_default(tmp_path):
    pdf = convert(make_docx(tmp_path / "a.docx", paragraphs=[ENGLISH], align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                            margins=(2.54, 6, 2.54, 6)))
    assert not any(l[4].rstrip().endswith("-") for l in lines(pdf))


def test_hyphenation_when_docx_enables_it(tmp_path):
    pdf = convert(make_docx(tmp_path / "a.docx", paragraphs=[ENGLISH], align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                            margins=(2.54, 6, 2.54, 6), hyphenate=True))
    assert any(l[4].rstrip().endswith("-") for l in lines(pdf))


def bottom_margin_text(pdf, margin_mm):
    with pymupdf.open(pdf) as doc:
        h = doc[0].rect.height
        return [l[4] for l in lines(pdf) if l[1] > h - margin_mm * PT_PER_MM]


def test_no_page_number_without_footer(tmp_path):
    pdf = convert(make_docx(tmp_path / "a.docx", margins=(2.54,) * 4))
    assert bottom_margin_text(pdf, 25.4) == []


def test_page_number_from_footer_field(tmp_path, capsys):
    pdf = convert(make_docx(tmp_path / "a.docx", margins=(2.54,) * 4, page_field=True))
    assert bottom_margin_text(pdf, 25.4) == ["1"]
    assert "tidak ikut ke PDF" not in capsys.readouterr().err


def test_footer_with_text_still_warned(tmp_path, capsys):
    convert(make_docx(tmp_path / "a.docx", footer="Rahasia perusahaan", page_field=True))
    assert "tidak ikut ke PDF: footer" in capsys.readouterr().err


# AC-6: font not installed
def test_missing_font_warns_and_uses_default(tmp_path, capsys):
    pdf = convert(make_docx(tmp_path / "a.docx", font="Font Yang Tidak Ada"))
    err = capsys.readouterr().err
    assert "Font Yang Tidak Ada" in err and "peringatan" in err
    assert all("Libertinus" in l[5] for l in lines(pdf))


# AC-7: missing or nonsense values fall back to the previous defaults
def test_missing_section_values_use_defaults(tmp_path):
    src = make_docx(tmp_path / "a.docx")
    d = docx.Document(src)
    sect = d.sections[0]._sectPr
    for tag in ("pgSz", "pgMar"):
        for el in sect.findall(f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{tag}"):
            sect.remove(el)
    d.save(src)
    pdf = convert(src)
    w, h = page_mm(pdf)
    assert w == pytest.approx(210, abs=1) and h == pytest.approx(297, abs=1)
    assert mm(min(l[0] for l in lines(pdf))) == pytest.approx(31.75, abs=0.5)


@pytest.mark.parametrize("margins", [(-2, 2, 2, 2), (2, 15, 2, 15)])
def test_nonsense_margins_use_defaults(tmp_path, margins):
    pdf = convert(make_docx(tmp_path / "a.docx", page=(21, 29.7), margins=margins))
    assert mm(min(l[0] for l in lines(pdf))) == pytest.approx(31.75, abs=0.5)


# AC-8: Markdown/TXT → PDF unchanged
def test_markdown_pdf_unchanged(tmp_path):
    src = tmp_path / "c.md"
    src.write_text("Paragraf satu " * 60, encoding="utf-8")
    pdf = convert(src)
    w, h = page_mm(pdf)
    assert w == pytest.approx(210, abs=1) and h == pytest.approx(297, abs=1)
    assert mm(min(l[0] for l in lines(pdf))) == pytest.approx(31.75, abs=0.5)
    body = [l for l in lines(pdf) if l[4] != "1"]  # pandoc's default template numbers pages
    assert {round(l[6]) for l in body} == {11}
    assert all("Libertinus" in l[5] for l in body)


# Formatting applied directly to paragraphs and runs (no styles), as many real documents do
def test_direct_formatting_is_followed(tmp_path):
    d = docx.Document()
    for text in (TEXT, TEXT):
        p = d.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        run.font.size = Pt(14)
        run.font.name = "New Computer Modern"
    d.save(tmp_path / "langsung.docx")
    pdf = convert(tmp_path / "langsung.docx")
    body = lines(pdf)
    assert {round(l[6]) for l in body} == {14}
    assert all("NewCM" in l[5] for l in body)
    assert body[1][3] - body[0][3] == pytest.approx(1.15 * 1.5 * 14, rel=0.05)


# docx_conf replaces pandoc's `conf`; a pandoc that stops using it must fail loudly
def test_unknown_template_fails(tmp_path):
    typ = tmp_path / "doc.typ"
    typ.write_text("#show: doc => template(doc)\nhalo", encoding="utf-8")
    with pytest.raises(ConvertError, match="pandoc"):
        check_conf(typ)
